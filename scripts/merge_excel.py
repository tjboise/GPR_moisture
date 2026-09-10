"""
Merge original Excel + Additional data (1).xlsx into one clean Excel file.
Output: GPR_moisture_merged.xlsx  (same 6-sheet format as original)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import numpy as np

EXCEL_ORIG = '../data/GPR measurement data in field.xlsx'
EXCEL_ADD  = '../data/Additional data (1).xlsx'
EXCEL_OUT  = '../data/GPR_moisture_merged.xlsx'

COND_MAP_ADD = {          # soil type label in additional file → condition index
    'Sand 2in':       0,
    'sand 4in':       1,
    'sandy clay 4in': 2,
}

# ── Read additional data ───────────────────────────────────────────────────
wb_add = openpyxl.load_workbook(EXCEL_ADD, data_only=True)
ws_add = wb_add['Sheet1']
rows_add = list(ws_add.iter_rows(values_only=True))

n_meas_add = 23
soil_types_add = [rows_add[0][4 + j] for j in range(n_meas_add)]

# Moisture: rows 2-5 → S, T, M, B
moist_add = {}          # cond_idx → list of [S, T, M, B] per sample
gpr_add   = {}          # cond_idx → list of 256-length arrays
for ci in range(3):
    moist_add[ci] = []
    gpr_add[ci]   = []

# Collect valid GPR rows (where time col is not None)
gpr_time_rows = [r for r in rows_add[9:] if r[3] is not None]  # 256 rows

for j in range(n_meas_add):
    stype  = soil_types_add[j]
    ci     = COND_MAP_ADD.get(stype, -1)
    if ci < 0:
        continue
    mo = [rows_add[ri][4 + j] for ri in [2, 3, 4, 5]]  # S, T, M, B
    mo = [float(v) if v is not None else None for v in mo]
    moist_add[ci].append(mo)
    trace = [float(r[4 + j]) if r[4 + j] is not None else 0.0
             for r in gpr_time_rows]
    gpr_add[ci].append(trace)

# Count how many new samples per condition
add_counts = {ci: len(gpr_add[ci]) for ci in range(3)}
print(f'Additional samples — 2in_sand: {add_counts[0]}, '
      f'4in_sand: {add_counts[1]}, 4in_clay: {add_counts[2]}')

# ── Style helpers ──────────────────────────────────────────────────────────
HDR_ORIG = PatternFill('solid', fgColor='DDEEFF')   # light blue — original
HDR_NEW  = PatternFill('solid', fgColor='FFEEDD')   # light orange — additional
BOLD     = Font(bold=True)

def style_header(cell, fill):
    cell.font      = BOLD
    cell.fill      = fill
    cell.alignment = Alignment(horizontal='center')

# ── Build merged workbook ──────────────────────────────────────────────────
wb_orig = openpyxl.load_workbook(EXCEL_ORIG, data_only=True)
wb_out  = openpyxl.Workbook()
wb_out.remove(wb_out.active)   # remove default empty sheet

SHEETS = [
    # (orig gpr sheet, orig moist sheet, cond_idx, short_name)
    ('2 in sand gpr data',   '2 in sand moisture data',  0, '2in_sand'),
    ('4in sand gpr data',    '4in sand mosture data',     1, '4in_sand'),
    ('4in clay gpr data',    '4in clay moisture data',    2, '4in_clay'),
]

for gpr_sname, moist_sname, ci, short in SHEETS:
    # ── GPR sheet ──────────────────────────────────────────────────────────
    ws_src = wb_orig[gpr_sname]
    orig_rows = list(ws_src.iter_rows(values_only=True))
    n_orig = ws_src.max_column - 1       # number of original traces
    n_new  = add_counts[ci]
    n_total = n_orig + n_new

    ws_gpr = wb_out.create_sheet(title=gpr_sname)

    # Header row (row 1): Time + No.1 … No.(n_total)
    ws_gpr.cell(1, 1, 'Time').font = BOLD
    for k in range(1, n_orig + 1):
        c = ws_gpr.cell(1, k + 1, f'No.{k}')
        style_header(c, HDR_ORIG)
    for k in range(n_new):
        c = ws_gpr.cell(1, n_orig + k + 2, f'No.{n_orig + k + 1}')
        style_header(c, HDR_NEW)

    # Data rows: original
    for ri, row in enumerate(orig_rows[1:], start=2):
        for ci2, val in enumerate(row):
            ws_gpr.cell(ri, ci2 + 1, val)

    # Data rows: additional (256 time steps)
    for ri, row in enumerate(orig_rows[1:], start=2):
        t_idx = ri - 2
        for k in range(n_new):
            ws_gpr.cell(ri, n_orig + k + 2,
                        gpr_add[ci][k][t_idx] if t_idx < 256 else None)

    # ── Moisture sheet ─────────────────────────────────────────────────────
    ws_msrc = wb_orig[moist_sname]
    moist_rows = list(ws_msrc.iter_rows(values_only=True))

    ws_mo = wb_out.create_sheet(title=moist_sname)

    # Header col labels from original (row 0 = label row in original)
    for ri, row in enumerate(moist_rows, start=1):
        for ci2, val in enumerate(row):
            ws_mo.cell(ri, ci2 + 1, val)

    # Determine which rows are S/T/M/B and avg in original
    # Original format: row0=header (No.1…), row1=S, row2=T, row3=M, row4=B, row5=avg(if present)
    # Find the label rows
    layer_rows = {}   # label_str → row_index (1-based in ws_mo)
    for ri, row in enumerate(moist_rows, start=1):
        lbl = str(row[0]).strip() if row[0] is not None else ''
        for layer in ['S', 'T', 'M', 'B']:
            if lbl.startswith(layer):
                layer_rows[layer] = ri
                break

    # Add column headers for new samples
    # Row 0 in original = column number headers
    for k in range(n_new):
        c = ws_mo.cell(1, n_orig + k + 2, f'No.{n_orig + k + 1}')
        style_header(c, HDR_NEW)

    # Add moisture values for new samples
    for layer_i, layer_key in enumerate(['S', 'T', 'M', 'B']):
        row_i = layer_rows.get(layer_key)
        if row_i is None:
            continue
        for k in range(n_new):
            ws_mo.cell(row_i, n_orig + k + 2, moist_add[ci][k][layer_i])

    # Recalculate average row if present
    # Find avg row (last non-empty row after B)
    avg_row = None
    for ri, row in enumerate(moist_rows, start=1):
        lbl = str(row[0]).strip().lower() if row[0] is not None else ''
        if 'avg' in lbl or 'average' in lbl:
            avg_row = ri
            break
    if avg_row is None and len(moist_rows) > max(layer_rows.values(), default=0):
        # Last row might be avg without label
        last_ri = len(moist_rows)
        if moist_rows[-1][0] is None and any(v is not None for v in moist_rows[-1][1:]):
            avg_row = last_ri

    if avg_row:
        layer_ri_list = list(layer_rows.values())
        for k in range(n_new):
            col = n_orig + k + 2
            vals = [ws_mo.cell(lr, col).value for lr in layer_ri_list
                    if ws_mo.cell(lr, col).value is not None]
            ws_mo.cell(avg_row, col, round(sum(vals) / len(vals), 2) if vals else None)

    print(f'  {short}: {n_orig} orig + {n_new} new = {n_total} total samples')

# ── Add a legend sheet ─────────────────────────────────────────────────────
ws_legend = wb_out.create_sheet(title='README', index=0)
info = [
    ['GPR Moisture Prediction — Merged Dataset'],
    [],
    ['Source files:'],
    ['  Original:', EXCEL_ORIG],
    ['  Additional:', EXCEL_ADD],
    [],
    ['Column color coding:'],
    ['  Blue header', '= original measurements'],
    ['  Orange header', '= additional measurements'],
    [],
    ['Moisture depth layers:'],
    ['  S', '0 cm (surface)'],
    ['  T', '8 cm'],
    ['  M', '22 cm'],
    ['  B', '35 cm'],
    [],
    ['Conditions:'],
    ['  2in_sand', 'Sand, 2-inch pipe, pipe top at 0.40 m depth'],
    ['  4in_sand', 'Sand, 4-inch pipe, pipe top at 0.35 m depth'],
    ['  4in_clay', 'Sandy clay, 4-inch pipe, pipe top at 0.30 m depth'],
    [],
    ['GPR signal:'],
    ['  256 time samples, dt = 0.099609 ns (~25.4 ns total window)'],
    ['  fs ≈ 10.04 GHz'],
]
for ri, row in enumerate(info, start=1):
    for ci2, val in enumerate(row):
        cell = ws_legend.cell(ri, ci2 + 1, val)
        if ri == 1:
            cell.font = Font(bold=True, size=13)
ws_legend.column_dimensions['A'].width = 22
ws_legend.column_dimensions['B'].width = 55

wb_out.save(EXCEL_OUT)
print(f'\nSaved: {EXCEL_OUT}')
